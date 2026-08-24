"""Persist JSON / CSV / Excel reports and evidence images for a crawl run."""

from __future__ import annotations

import csv
import json
import os
from datetime import datetime, timezone
from typing import List

REPORTS_DIR = os.path.join(os.path.dirname(__file__), "reports")
ISSUE_FIELDS = ["Page", "Element", "Issue", "Confidence", "Status", "Details", "EvidenceImage"]


def ensure_run_dir(run_id: str) -> str:
    path = os.path.join(REPORTS_DIR, run_id)
    os.makedirs(path, exist_ok=True)
    os.makedirs(os.path.join(path, "evidence"), exist_ok=True)
    return path


def strip_internal(issue: dict) -> dict:
    return {k: issue.get(k, "") for k in ISSUE_FIELDS}


def write_json(run_dir: str, payload: dict) -> str:
    path = os.path.join(run_dir, "report.json")
    serializable = json.loads(json.dumps(payload, default=str))
    with open(path, "w", encoding="utf-8") as f:
        json.dump(serializable, f, indent=2, ensure_ascii=False)
    return path


def write_csv(run_dir: str, issues: List[dict]) -> str:
    path = os.path.join(run_dir, "report.csv")
    with open(path, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=ISSUE_FIELDS, extrasaction="ignore")
        writer.writeheader()
        for issue in issues:
            row = strip_internal(issue)
            evidence = row.get("EvidenceImage") or ""
            if evidence.startswith("data:image"):
                row["EvidenceImage"] = "(embedded data URL; see JSON/evidence folder)"
            writer.writerow(row)
    return path


def write_excel(run_dir: str, payload: dict, issues: List[dict]) -> str:
    path = os.path.join(run_dir, "report.xlsx")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = "UI Validation Framework Report"
    summary["A1"].font = Font(bold=True, size=14)
    meta = payload.get("summary") or {}
    rows = [
        ("Generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
        ("Root URL", payload.get("root_url", "")),
        ("Baseline Root URL", payload.get("baseline_root_url") or "(none)"),
        ("Pages crawled", meta.get("pages_crawled", 0)),
        ("FAIL issues", meta.get("fail_count", 0)),
        ("PASS rows", meta.get("pass_count", 0)),
        ("Run ID", payload.get("run_id", "")),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        summary[f"A{i}"] = k
        summary[f"B{i}"] = v

    ws = wb.create_sheet("Issues")
    header_fill = PatternFill("solid", fgColor="0696D7")
    header_font = Font(bold=True, color="FFFFFF")
    for col, name in enumerate(ISSUE_FIELDS, start=1):
        cell = ws.cell(1, col, name)
        cell.fill = header_fill
        cell.font = header_font

    fail_fill = PatternFill("solid", fgColor="FECACA")
    pass_fill = PatternFill("solid", fgColor="BBF7D0")
    for r, issue in enumerate(issues, start=2):
        row = strip_internal(issue)
        evidence = row.get("EvidenceImage") or ""
        if isinstance(evidence, str) and evidence.startswith("data:image"):
            row["EvidenceImage"] = "(see evidence/ folder or JSON)"
        for c, name in enumerate(ISSUE_FIELDS, start=1):
            cell = ws.cell(r, c, row.get(name, ""))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        fill = fail_fill if row.get("Status") == "FAIL" else pass_fill
        ws.cell(r, 5).fill = fill

    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["F"].width = 60
    ws.column_dimensions["G"].width = 36

    pages = wb.create_sheet("Pages")
    pages.append(["URL", "Title", "Depth", "Status", "Elements", "Error"])
    for page in payload.get("pages") or []:
        pages.append([
            page.get("url"),
            page.get("title"),
            page.get("depth"),
            page.get("http_status"),
            (page.get("element_counts") or {}).get("total"),
            page.get("error") or "",
        ])

    wb.save(path)
    return path


def save_png(run_dir: str, name: str, png_bytes: bytes) -> str:
    rel = os.path.join("evidence", name)
    path = os.path.join(run_dir, rel)
    with open(path, "wb") as f:
        f.write(png_bytes)
    return rel.replace("\\", "/")
